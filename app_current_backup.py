import os
import uuid
import base64
from flask import Flask, request, render_template, redirect, url_for, flash, send_file
from werkzeug.utils import secure_filename
from pptx import Presentation
from pptx.exc import PackageNotFoundError
from zipfile import BadZipFile
from docx import Document

app = Flask(__name__)
app.secret_key = "b5f8c27a6c7a4e8d9f561c6277e739bc"  # Replace with a secure key

# Ensure uploads folder exists
UPLOAD_FOLDER = "uploads"
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER

# Ensure slide images folder exists
os.makedirs("static/slide_images", exist_ok=True)

# Allowed file extensions
ALLOWED_EXTENSIONS = {"pptx"}


def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


@app.route("/", methods=["GET", "POST"])
def upload_file():
    if request.method == "POST":
        if "file" not in request.files:
            flash("No file part")
            return redirect(url_for("upload_file"))

        file = request.files["file"]
        if file.filename == "":
            flash("No selected file")
            return redirect(request.url)

        if file and allowed_file(file.filename):
            filename = secure_filename(file.filename)
            filepath = os.path.join(app.config["UPLOAD_FOLDER"], filename)
            file.save(filepath)

            try:
                slides_data = parse_pptx(filepath)
                output_file = convert_to_word(slides_data)
                return redirect(url_for("download_file", filename=output_file))
            except Exception as e:
                flash(f"An error occurred: {str(e)}")
                return redirect(request.url)

    return render_template("index.html")


# Function to parse PowerPoint file
def parse_pptx(filepath):
    try:
        prs = Presentation(filepath)
    except (PackageNotFoundError, BadZipFile):
        raise Exception("Invalid or corrupted PowerPoint file.")

    slides_data = []

    for i, slide in enumerate(prs.slides):
        slide_num = i + 1
        title = None
        text_html = ""
        table_html = ""

        # Extract title
        for shape in slide.shapes:
            if shape.is_placeholder and shape.placeholder_format.type == 1:  # TITLE
                title = shape.text

            # Extract text
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text_html += paragraph.text + "\n"

            # Extract tables (if any)
            if shape.has_table:
                table = shape.table
                for row in table.rows:
                    table_html += "\t".join([cell.text for cell in row.cells]) + "\n"

        slides_data.append({
            "title": title or f"Slide {slide_num}",
            "text_html": text_html.strip(),
            "table_html": table_html.strip(),
        })

    return slides_data


# Function to convert slides data to Word
def convert_to_word(slides_data):
    unique_id = str(uuid.uuid4())
    output_path = os.path.join(app.config["UPLOAD_FOLDER"], f"output_{unique_id}.docx")
    doc = Document()

    for slide in slides_data:
        # Add slide title
        if slide["title"]:
            doc.add_heading(slide["title"], level=1)

        # Add slide text
        if slide["text_html"]:
            doc.add_paragraph(slide["text_html"])

        # Add slide table (if any)
        if slide["table_html"]:
            doc.add_paragraph("Table:")
            table = doc.add_table(rows=1, cols=len(slide["table_html"].split("\t")))
            table.style = "Table Grid"
            first_row = table.rows[0].cells
            for i, cell in enumerate(slide["table_html"].split("\n")[0].split("\t")):
                first_row[i].text = cell

            for row_data in slide["table_html"].split("\n")[1:]:
                row = table.add_row().cells
                for i, cell_data in enumerate(row_data.split("\t")):
                    row[i].text = cell_data

        doc.add_paragraph("\n")

    doc.save(output_path)
    return f"output_{unique_id}.docx"


# Route to handle file download
@app.route("/download/<filename>")
def download_file(filename):
    file_path = os.path.join(app.config["UPLOAD_FOLDER"], filename)
    if not os.path.exists(file_path):
        flash("File not found.")
        return redirect(url_for("upload_file"))
    return send_file(file_path, as_attachment=True)


if __name__ == "__main__":
    app.run(debug=True, host="0.0.0.0", port=5001)

